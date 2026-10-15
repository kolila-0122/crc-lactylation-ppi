# -*- coding: utf-8 -*-
"""manuscript.md -> manuscript.docx（投稿版 Word 正文）"""
import os, re
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# 包根目录：脚本上级（投稿包内 code/ 的上级）；可用 PACKAGE_DIR 环境变量覆盖
SRC = os.environ.get("PACKAGE_DIR", os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
MD = os.path.join(SRC, "manuscript.md")
DOCX = os.path.join(SRC, "manuscript.docx")

doc = Document()

# 基础样式：正文 Times New Roman 12pt
style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(12)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")  # 防中文乱码（正文无中文，保险起见）

def set_run_font(run, size=12, bold=False, italic=False, color=(0, 0, 0)):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(*color)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

def add_heading(text, level):
    sizes = {1: 16, 2: 14, 3: 12.5}
    p = doc.add_paragraph()
    p.space_before = Pt(12 if level == 1 else 8)
    p.space_after = Pt(6)
    r = p.add_run(text)
    set_run_font(r, size=sizes.get(level, 12), bold=True)
    return p

def add_para(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    # 处理 **bold** 和 *italic* 内联标记
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*)', text)
    for t in tokens:
        if not t:
            continue
        if t.startswith("**") and t.endswith("**"):
            r = p.add_run(t[2:-2]); set_run_font(r, bold=True)
        elif t.startswith("*") and t.endswith("*") and len(t) > 2:
            r = p.add_run(t[1:-1]); set_run_font(r, italic=True)
        else:
            r = p.add_run(t); set_run_font(r)
    return p

def set_three_line(t):
    """三线表：仅顶部、表头下、底部有横线，无竖线"""
    tbl = t._tbl
    tblPr = tbl.tblPr
    # 移除默认样式边框，改自定义
    for old in tblPr.findall(qn('w:tblBorders')):
        tblPr.remove(old)
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:%s' % edge)
        el.set(qn('w:val'), 'none'); el.set(qn('w:sz'), '0'); el.set(qn('w:space'), '0')
        borders.append(el)
    borders.find(qn('w:top')).set(qn('w:val'), 'single'); borders.find(qn('w:top')).set(qn('w:sz'), '12')
    borders.find(qn('w:bottom')).set(qn('w:val'), 'single'); borders.find(qn('w:bottom')).set(qn('w:sz'), '12')
    tblPr.append(borders)
    # 表头行下边框
    for cell in t.rows[0].cells:
        tcPr = cell._tc.get_or_add_tcPr()
        tcB = OxmlElement('w:tcBorders')
        btm = OxmlElement('w:bottom')
        btm.set(qn('w:val'), 'single'); btm.set(qn('w:sz'), '6'); btm.set(qn('w:space'), '0')
        tcB.append(btm)
        tcPr.append(tcB)

def add_table(rows):
    # rows: list of list of str；rows[0] 表头，rows[1] 分隔行（跳过），rows[2:] 数据
    header = rows[0]
    body = rows[2:]
    ncol = len(header)
    t = doc.add_table(rows=1, cols=ncol)
    # 表头
    for j, cell_text in enumerate(header):
        cell = t.rows[0].cells[j]
        cell.text = ""
        r = cell.paragraphs[0].add_run(cell_text.strip())
        set_run_font(r, size=10.5, bold=True)
    # 数据
    for row in body:
        cells = t.add_row().cells
        for j in range(ncol):
            val = row[j].strip() if j < len(row) else ""
            cells[j].text = ""
            r = cells[j].paragraphs[0].add_run(val)
            set_run_font(r, size=10.5)
    set_three_line(t)
    doc.add_paragraph()  # 表格后空行
    return t

# 逐行解析
lines = open(MD, encoding="utf-8").read().split("\n")
i = 0
table_buf = []
in_table = False
while i < len(lines):
    line = lines[i].rstrip()
    s = line.strip()
    # 表格
    if s.startswith("|") and s.endswith("|"):
        if not in_table:
            in_table = True
            table_buf = []
        table_buf.append([c.strip() for c in s.strip("|").split("|")])
        i += 1
        continue
    else:
        if in_table:
            add_table(table_buf)
            in_table = False
            table_buf = []
    # 标题
    if s.startswith("### "):
        add_heading(s[4:], 3)
    elif s.startswith("## "):
        add_heading(s[3:], 2)
    elif s.startswith("# "):
        add_heading(s[2:], 1)
    elif s.startswith("- "):
        add_para("• " + s[2:])
    elif s.startswith("|") is False and s == "":
        pass
    elif re.match(r'^\d+\.\s', s):  # 参考文献编号列表
        add_para(s)
    elif s.startswith("**") and s.endswith("**"):
        add_para(s)
    elif s:
        add_para(s)
    i += 1
if in_table:
    add_table(table_buf)

doc.save(DOCX)
print("DOCX saved:", DOCX)

# 机检：表格数
from docx import Document as D2
chk = D2(DOCX)
print("tables in docx:", len(chk.tables))
print("paragraphs:", len(chk.paragraphs))
